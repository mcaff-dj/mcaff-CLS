#!/usr/bin/env python3
"""One-off backfill: marks tickets Delivered in PEP_CLS.Delivery_escalation from
Mapping_delivery.xlsx (columns: AWB Number, Delivered Date) - sets outcome='Delivered' and
disposed_at=<that AWB's Delivered Date>. tat_bucket isn't a stored column (see
api/_lib/db.js's getDeliveryEscalationHistory) - it's computed at query time from
disposed_at minus added_date, so backfilling just these two columns is enough for TAT to
fall out correctly on the next read.

Only touches Fresh-eligible rows (outcome blank/RTO/Escalated - same set
getDeliveryEscalationFresh lists, same guard api/delivery-escalation/record.js's bulkDispose
enforces), so an already-Delivered or otherwise-disposed row's history is never overwritten.
Matches EVERY table row sharing an AWB (same "update all matches" policy as the app's own
bulk upload) - an AWB can legitimately repeat across brands or a re-shipped order.

Chunked IN(...) + CASE update, not a staging-table JOIN: this DB user has no CREATE TEMPORARY
TABLES grant (confirmed live - "Access denied ... to database 'PEP_CLS'" on the CREATE
TEMPORARY TABLE itself, despite being able to SELECT/INSERT/UPDATE/ALTER it fine elsewhere).
Each chunk (CHUNK_SIZE AWBs) becomes one UPDATE with a CASE awb_code WHEN...THEN... per row and
an IN(...) filter - still server-side matching, just without needing any table-creation
privilege at all.

Does not touch agent_email/assigned_at/agent_remarks - nobody actually worked these
tickets, so there's no agent or remarks to backfill.

Dry-run by default (reports match counts, no write); --apply performs the UPDATE.
"""
import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import mysql_lib

MAPPING_FILE = Path(__file__).resolve().parent.parent / "Mapping_delivery.xlsx"
CHUNK_SIZE = 500

# %% (not %), because every execute() call below this uses it alongside a params tuple -
# pymysql only %-substitutes when args is passed, and a bare '%' in that mode would be
# misread as a format directive.
FRESH_ELIGIBLE_SQL = """(d.outcome IS NULL OR d.outcome = ''
     OR d.outcome = 'RTO' OR d.outcome LIKE 'RTO > %%'
     OR d.outcome = 'Escalated' OR d.outcome LIKE 'Escalated > %%')"""


def dedupe_by_awb(rows):
    """Last occurrence wins on a duplicate AWB within the file itself - keeps each chunk's
    CASE deterministic. Rows with a blank AWB or Delivered Date are dropped."""
    by_awb = {}
    for awb, delivered_date in rows:
        if awb is None or delivered_date is None:
            continue
        by_awb[str(awb).strip()] = delivered_date
    return list(by_awb.items())


def load_mapping(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    return dedupe_by_awb(ws.iter_rows(min_row=2, values_only=True))


def chunk(seq, size):
    for i in range(0, len(seq), size):
        yield seq[i:i + size]


def count_chunk(cur, batch):
    awbs = [awb for awb, _ in batch]
    placeholders = ", ".join(["%s"] * len(awbs))
    cur.execute(
        f"SELECT COUNT(*), SUM({FRESH_ELIGIBLE_SQL}) FROM Delivery_escalation d WHERE d.awb_code IN ({placeholders})",
        awbs,
    )
    found, eligible = cur.fetchone()
    return found or 0, int(eligible or 0)


def update_chunk(cur, batch):
    case_when = " ".join(["WHEN %s THEN %s"] * len(batch))
    awbs = [awb for awb, _ in batch]
    placeholders = ", ".join(["%s"] * len(awbs))
    case_params = [v for pair in batch for v in pair]
    cur.execute(
        f"""
        UPDATE Delivery_escalation d
        SET d.outcome = 'Delivered',
            d.disposed_at = CASE d.awb_code {case_when} END
        WHERE d.awb_code IN ({placeholders}) AND {FRESH_ELIGIBLE_SQL}
        """,
        case_params + awbs,
    )
    return cur.rowcount


def self_check():
    from datetime import date
    rows = [("123", date(2026, 1, 1)), ("123", date(2026, 1, 2)), (None, date(2026, 1, 1)), ("456", None)]
    assert dedupe_by_awb(rows) == [("123", date(2026, 1, 2))], dedupe_by_awb(rows)
    assert list(chunk([1, 2, 3, 4, 5], 2)) == [[1, 2], [3, 4], [5]]
    print("self-check ok")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="Perform the update (default is a dry run).")
    ap.add_argument("--file", default=str(MAPPING_FILE))
    ap.add_argument("--self-check", action="store_true")
    args = ap.parse_args()
    if args.self_check:
        return self_check()

    mapping = load_mapping(args.file)
    print(f"{len(mapping)} unique AWB(s) with a Delivered Date in {args.file}")
    if not mapping:
        return

    cred = mysql_lib.get_credential()
    if cred is None:
        raise SystemExit("MYSQL_* credentials not configured.")
    import pymysql
    conn = pymysql.connect(
        host=cred["host"], user=cred["user"], password=cred["password"],
        database="PEP_CLS", port=cred["port"], ssl={"ssl": {}}, connect_timeout=15,
    )
    try:
        cur = conn.cursor()
        batches = list(chunk(mapping, CHUNK_SIZE))

        total_found = total_eligible = total_updated = 0
        for i, batch in enumerate(batches, 1):
            found, eligible = count_chunk(cur, batch)
            total_found += found
            total_eligible += eligible
            updated = update_chunk(cur, batch) if (args.apply and eligible) else 0
            total_updated += updated
            suffix = f", {updated} updated" if args.apply else ""
            print(f"  batch {i}/{len(batches)}: {found} matched, {eligible} Fresh-eligible{suffix}")

        print(f"{total_found} Delivery_escalation row(s) match an AWB in the file")
        print(f"{total_eligible} of those are Fresh-eligible (blank/RTO/Escalated) - only these get updated")
        print(f"{total_found - total_eligible} already resolved some other way - left untouched")

        if not args.apply:
            print("Dry run - re-run with --apply to perform the update.")
            return

        conn.commit()
        print(f"updated {total_updated} row(s)")
    finally:
        conn.close()


if __name__ == "__main__":
    main()
